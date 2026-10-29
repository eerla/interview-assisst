"""
Interview Assistant - Main Application
Streamlit app for AI-powered interview question generation
"""

import io
import json
import streamlit as st
import pandas as pd  # type: ignore[import-untyped]
from docx import Document
from reportlab.lib.pagesizes import letter  # type: ignore[import-untyped]
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # type: ignore[import-untyped]
from reportlab.lib.units import inch  # type: ignore[import-untyped]
from reportlab.platypus import (  # type: ignore[import-untyped]
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib import colors  # type: ignore[import-untyped]
from typing import Any, Dict, List, Optional

from interview_assistant.core.resume_parser import ResumeParser
from interview_assistant.core.question_generator import QuestionGenerator
from interview_assistant.utils.helpers import (
    format_questions_for_export,
    validate_file_upload,
)

# Page configuration
st.set_page_config(
    page_title="Interview Assistant",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Custom CSS for better styling
st.markdown(
    """
<style>
    .main-header {
        font-size: 3rem;
        font-weight: bold;
        color: #1f77b4;
        text-align: center;
        margin-bottom: 2rem;
    }
    .sub-header {
        font-size: 1.5rem;
        color: #666;
        text-align: center;
        margin-bottom: 2rem;
    }
    .upload-section {
        background-color: #f8f9fa;
        padding: 2rem;
        border-radius: 10px;
        border: 2px dashed #dee2e6;
        text-align: center;
        margin: 2rem 0;
    }
    .question-card {
        background-color: black;
        padding: 0.5rem;
        border-radius: 8px;
        border-left: 4px solid #1f77b4;
        margin: 1rem 0;
        color: white;
        box-shadow: 0 2px 4px rgba(0,0,0,0.1);
    }
    .difficulty-easy { border-left-color: #28a745; }
    .difficulty-medium { border-left-color: #ffc107; }
    .difficulty-hard { border-left-color: #dc3545; }
    .category-badge {
        background-color: #e9ecef;
        color: #495057;
        padding: 0.25rem 0.75rem;
        border-radius: 15px;
        font-size: 0.875rem;
        font-weight: 500;
    }
    .category-badge.easy {
        background-color: #d4edda !important;
        color: #155724 !important;
    }
    .category-badge.medium {
        background-color: #fff3cd !important;
        color: #856404 !important;
    }
    .category-badge.hard {
        background-color: #f8d7da !important;
        color: #721c24 !important;
    }
</style>
""",
    unsafe_allow_html=True,
)

QUESTIONS_PER_PAGE = 5


def main() -> None:
    """Main application function"""

    # Header
    st.markdown(
        '<h1 class="main-header">🎯 Interview Assistant</h1>', unsafe_allow_html=True
    )
    st.markdown(
        '<p class="sub-header">AI-powered interview question generator for technical roles</p>',
        unsafe_allow_html=True,
    )

    # Sidebar for configuration
    with st.sidebar:
        st.header("⚙️ Configuration")

        # Question count slider
        question_count = st.slider(
            "Number of Questions",
            min_value=1,
            max_value=10,
            value=5,
            help="Select how many questions to generate (maximum 10)",
        )

        # Difficulty filter
        difficulty_filter = st.multiselect(
            "Difficulty Level",
            options=["Easy", "Medium", "Hard"],
            default=["Easy"],
            help="Select difficulty levels for questions (Easy: Basic knowledge, Medium: Practical experience, Hard: Advanced concepts)",
        )

        # Category filter
        category_filter = st.multiselect(
            "Question Categories",
            options=[
                "Technical Skills",
                "Experience & Projects",
                "Problem Solving",
                "Behavioral",
                "Coding Test",
            ],
            default=["Technical Skills"],
            help="Select categories to focus on (Technical: Skills & tools, Experience: Past work, Problem Solving: Analytical thinking, Behavioral: Soft skills, Coding Test: LeetCode-style coding questions)",
        )

        # Validation
        if not difficulty_filter:
            st.error("⚠️ Please select at least one difficulty level")
        if not category_filter:
            st.error("⚠️ Please select at least one category")

        st.markdown("---")
        st.markdown("### 📋 Instructions")
        st.markdown("""
        1. Upload a resume (PDF or Word format)
        2. Select number of questions (max 10)
        3. Choose difficulty levels and categories
        4. Click 'Generate Questions' button
        5. Review and export the generated questions
        """)

    # Main content area
    col1, col2 = st.columns([1, 0.5])

    with col1:
        st.header("📄 Resume Upload")

        # File upload section
        uploaded_file = st.file_uploader(
            "Choose a resume file",
            type=["pdf", "docx"],
            help="Upload a PDF or Word document containing the candidate's resume",
        )

        if uploaded_file is not None:
            # Validate file upload
            if not validate_file_upload(uploaded_file):
                return

            st.success(f"✅ File uploaded: {uploaded_file.name}")

            # Display file info
            file_details = {
                "Filename": uploaded_file.name,
                "File type": uploaded_file.type,
                "File size": f"{uploaded_file.size / 1024:.1f} KB",
            }
            st.json(file_details)

            # Parse resume
            with st.spinner("📖 Parsing resume..."):
                resume_text = ResumeParser.parse_resume(uploaded_file)

            if not resume_text:
                st.error(
                    "❌ Failed to extract text from the resume. Please try a different file."
                )
                return

            st.success("✅ Resume parsed successfully!")

            # Show resume preview
            with st.expander("📋 Resume Preview"):
                # Show cleaned, formatted text
                cleaned_preview = (
                    resume_text[:1000] + "..."
                    if len(resume_text) > 1000
                    else resume_text
                )
                st.text_area(
                    "Cleaned Resume Content",
                    cleaned_preview,
                    height=200,
                    help="This preview is truncated for display only; the full resume is used for question generation",
                )

                # Show text statistics
                col1, col2, col3 = st.columns(3)
                with col1:
                    st.metric("Total Characters", len(resume_text))
                with col2:
                    st.metric("Words", len(resume_text.split()))
                with col3:
                    st.metric("Lines", resume_text.count("\n") + 1)

            # Generate questions button
            if st.button(
                "🚀 Generate Questions", type="primary", use_container_width=True
            ):
                if not difficulty_filter:
                    st.error("⚠️ Please select at least one difficulty level")
                elif not category_filter:
                    st.error("⚠️ Please select at least one category")
                elif question_count > 10:
                    st.error("⚠️ Question count cannot exceed 10")
                else:
                    st.session_state.difficulty_filter = difficulty_filter
                    st.session_state.category_filter = category_filter
                    generate_questions(
                        resume_text,
                        question_count,
                        difficulty_filter,
                        category_filter,
                    )

            # Display generated questions below the button
            if "questions" in st.session_state and st.session_state.questions:
                st.markdown("---")
                st.subheader("🎯 Generated Questions")
                col_q, col_clear = st.columns([4, 1])
                with col_clear:
                    if st.button("🗑️ Clear Results", use_container_width=True):
                        for key in [
                            "questions",
                            "resume_insights",
                            "ats_suggestions",
                            "resume_text",
                            "difficulty_filter",
                            "category_filter",
                            "current_page",
                            "answers",
                            "evaluations",
                            "edit_mode",
                        ]:
                            st.session_state.pop(key, None)
                        st.rerun()
                with col_q:
                    display_questions(
                        st.session_state.questions,
                        st.session_state.get("difficulty_filter", []),
                        st.session_state.get("category_filter", []),
                    )

    with col2:
        st.header("📊 Question Statistics")

        # Show summary info
        if "questions" in st.session_state and st.session_state.questions:
            questions = st.session_state.questions
            categories: Dict[str, List[Dict[str, Any]]] = {}
            difficulties: Dict[str, int] = {"Easy": 0, "Medium": 0, "Hard": 0}
            for question in questions:
                category = question["category"]
                if category not in categories:
                    categories[category] = []
                categories[category].append(question)
                diff = question.get("difficulty", "Unknown")
                difficulties[diff] = difficulties.get(diff, 0) + 1

            # Show statistics
            col1, col2, col3 = st.columns(3)
            with col1:
                st.metric("Total Questions", len(questions))
            with col2:
                st.metric("Categories", len(categories))
            with col3:
                st.metric("Hard Questions", difficulties.get("Hard", 0))

            # Show difficulty breakdown
            st.markdown("---")
            st.subheader("🎚️ Difficulty Breakdown")
            diff_col1, diff_col2, diff_col3 = st.columns(3)
            with diff_col1:
                st.metric("Easy", difficulties.get("Easy", 0))
            with diff_col2:
                st.metric("Medium", difficulties.get("Medium", 0))
            with diff_col3:
                st.metric("Hard", difficulties.get("Hard", 0))

            # Show category breakdown
            st.markdown("---")
            st.subheader("📂 Category Breakdown")
            for category, category_questions in categories.items():
                st.write(f"**{category}**: {len(category_questions)} questions")
        else:
            st.info(
                "👆 Upload a resume and click 'Generate Questions' to see statistics!"
            )


def generate_questions(
    resume_text: str,
    question_count: int,
    difficulty_filter: List[str],
    category_filter: List[str],
) -> None:
    """Generate interview questions using OpenAI"""
    generator = QuestionGenerator()
    try:
        # Sanitize the full resume text before sending to the model
        sanitized_text = generator._sanitize_resume_text(resume_text)

        with st.spinner("🤖 Generating questions with AI..."):
            response = generator.generate_questions(
                sanitized_text, question_count, difficulty_filter, category_filter
            )

            # Extract data from response
            questions = response.get("questions", [])
            insights = response.get("insights", {})
            ats_suggestions = response.get("ats_suggestions", [])

            # Store in session state
            st.session_state.questions = questions
            st.session_state.resume_insights = insights
            st.session_state.ats_suggestions = ats_suggestions
            st.session_state.resume_text = sanitized_text
            st.session_state.current_page = 0
            # Initialize answer storage
            st.session_state.answers = {i: "" for i in range(len(questions))}
            st.session_state.evaluations = {}
            st.session_state.edit_mode = {}
            st.success(f"✅ Generated {len(questions)} questions!")

    except RuntimeError as e:
        st.error(f"❌ {e}")
    except Exception as e:
        st.error(f"❌ Failed to generate questions: {str(e)}")


def _global_question_index(
    filtered_questions: List[Dict[str, Any]], question: Dict[str, Any]
) -> int:
    """Find the global index of a question within the filtered list."""
    for idx, q in enumerate(filtered_questions):
        if q is question:
            return idx
    return 0


def regenerate_single_question(
    question: Dict[str, Any],
    difficulty_filter: List[str],
    category_filter: List[str],
    q_idx: int,
    filtered_questions: List[Dict[str, Any]],
) -> None:
    """Regenerate a single question using the LLM."""
    generator = QuestionGenerator()
    try:
        resume_text = st.session_state.get("resume_text", "")
        prompt = generator._create_question_prompt(
            resume_text,
            1,
            [question.get("difficulty", "Medium")],
            [question.get("category", "Technical Skills")],
        )
        response = generator._create_chat_completion(
            model=generator.model,
            messages=[
                {
                    "role": "system",
                    "content": generator._build_system_prompt(resume_text),
                },
                {"role": "user", "content": prompt},
            ],
            max_completion_tokens=1000,
            response_format={"type": "json_object"},
        )
        llm_response = response.choices[0].message.content
        if not llm_response or llm_response.strip() == "":
            st.error("❌ Failed to regenerate question: empty response")
            return

        cleaned = generator._strip_markdown_fences(llm_response)
        data = json.loads(cleaned)
        new_questions = data.get("questions", [])
        if new_questions:
            new_q = new_questions[0]
            new_q["category"] = question.get("category", "Technical Skills")
            new_q["difficulty"] = question.get("difficulty", "Medium")
            filtered_questions[q_idx] = new_q
            st.session_state.questions = filtered_questions
            st.session_state.answers[q_idx] = ""
            if q_idx in st.session_state.get("evaluations", {}):
                del st.session_state.evaluations[q_idx]
            st.success("✅ Question regenerated!")
        else:
            st.error("❌ Failed to regenerate question: no questions returned")
    except Exception as e:
        st.error(f"❌ Failed to regenerate question: {str(e)}")


def evaluate_single_question(
    question: Dict[str, Any],
    answer: str,
    q_idx: int,
) -> None:
    """Evaluate a single question answer and store result in session state."""
    generator = QuestionGenerator()
    try:
        resume_text = st.session_state.get("resume_text", "")
        result = generator.evaluate_answer(
            question=question["question"],
            category=question.get("category", "Technical Skills"),
            difficulty=question.get("difficulty", "Medium"),
            candidate_answer=answer,
            resume_context=resume_text,
        )
        if "evaluations" not in st.session_state:
            st.session_state["evaluations"] = {}
        st.session_state["evaluations"][q_idx] = result
    except Exception as e:
        st.error(f"❌ Failed to evaluate answer: {str(e)}")


def display_questions(
    questions: List[Dict[str, Any]],
    difficulty_filter: Optional[List[str]] = None,
    category_filter: Optional[List[str]] = None,
) -> None:
    """Display generated questions with filtering"""

    filtered_questions = questions
    if difficulty_filter:
        filtered_questions = [
            q for q in filtered_questions if q.get("difficulty") in difficulty_filter
        ]
    if category_filter:
        filtered_questions = [
            q for q in filtered_questions if q.get("category") in category_filter
        ]
    if not filtered_questions:
        st.warning(
            "No questions match the current filters. Try adjusting the filter settings."
        )
        return

    # Pagination
    total_questions = len(filtered_questions)
    total_pages = (total_questions + QUESTIONS_PER_PAGE - 1) // QUESTIONS_PER_PAGE

    current_page = st.session_state.get("current_page", 0)
    if current_page < 0 or current_page >= total_pages:
        current_page = 0
        st.session_state.current_page = 0

    start_idx = current_page * QUESTIONS_PER_PAGE
    end_idx = min(start_idx + QUESTIONS_PER_PAGE, total_questions)
    page_questions = filtered_questions[start_idx:end_idx]

    # Pagination controls
    if total_pages > 1:
        col_prev, col_page, col_next = st.columns([1, 2, 1])
        with col_prev:
            if st.button(
                "◀ Previous", disabled=(current_page == 0), use_container_width=True
            ):
                st.session_state.current_page = current_page - 1
                st.rerun()
        with col_page:
            st.write(f"Page {current_page + 1} of {total_pages}")
        with col_next:
            if st.button(
                "Next ▶",
                disabled=(current_page >= total_pages - 1),
                use_container_width=True,
            ):
                st.session_state.current_page = current_page + 1
                st.rerun()
        st.markdown("---")

    # Display questions by category
    categories: Dict[str, List[Dict[str, Any]]] = {}
    for question in page_questions:
        category = question["category"]
        if category not in categories:
            categories[category] = []
        categories[category].append(question)

    for category, category_questions in categories.items():
        st.subheader(f"📂 {category}")
        for i, question in enumerate(category_questions, 1):
            q_idx = _global_question_index(filtered_questions, question)

            difficulty_class = f"difficulty-{question['difficulty'].lower()}"
            difficulty_badge_class = f"category-badge {question['difficulty'].lower()}"

            # Question header with edit/regenerate actions
            col_q, col_actions = st.columns([6, 1])
            with col_q:
                if category == "Coding Test":
                    st.markdown(
                        f"""
                    <div class="question-card {difficulty_class}">
                        <div style="display: flex; justify-content: end; align-items: center; margin-bottom: 0.5rem;">
                            <span class="{difficulty_badge_class}">{question['difficulty']}</span>
                        </div>
                    </div>
                    """,
                        unsafe_allow_html=True,
                    )
                    st.markdown(f"**{q_idx}. {question['question']}**")
                    if "instructions" in question:
                        st.markdown(question["instructions"])
                    if "test_cases" in question:
                        st.markdown("**Sample Test Cases:**")
                        for tc in question["test_cases"]:
                            st.code(tc, language="python")
                else:
                    st.markdown(
                        f"""
                    <div class="question-card {difficulty_class}">
                        <div style="display: flex; justify-content: end; align-items: center; margin-bottom: 0.5rem;">
                            <span class="{difficulty_badge_class}">
                                {question['difficulty']}
                            </span>
                        </div>
                        <p style="margin: 0; font-size: 1.1rem;"><strong>{q_idx}.</strong> {question['question']}</p>
                    </div>
                    """,
                        unsafe_allow_html=True,
                    )

            with col_actions:
                edit_key = f"edit_{q_idx}"
                if st.button("✏️ Edit", key=edit_key, use_container_width=True):
                    edit_mode = st.session_state.get("edit_mode", {})
                    edit_mode[q_idx] = not edit_mode.get(q_idx, False)
                    st.session_state.edit_mode = edit_mode
                    st.rerun()

            # Edit mode UI
            edit_mode = st.session_state.get("edit_mode", {})
            if edit_mode.get(q_idx, False):
                edited_question = st.text_area(
                    "Edit question:",
                    value=question["question"],
                    key=f"edit_text_{q_idx}",
                    height=100,
                )
                col_save, col_regen, col_cancel = st.columns(3)
                with col_save:
                    if st.button(
                        "💾 Save", key=f"save_{q_idx}", use_container_width=True
                    ):
                        filtered_questions[q_idx]["question"] = edited_question
                        edit_mode = st.session_state.get("edit_mode", {})
                        edit_mode[q_idx] = False
                        st.session_state.edit_mode = edit_mode
                        st.rerun()
                with col_regen:
                    if st.button(
                        "🔄 Regenerate",
                        key=f"regen_{q_idx}",
                        use_container_width=True,
                    ):
                        regenerate_single_question(
                            question,
                            difficulty_filter or ["Medium"],
                            category_filter or ["Technical Skills"],
                            q_idx,
                            filtered_questions,
                        )
                        edit_mode = st.session_state.get("edit_mode", {})
                        edit_mode[q_idx] = False
                        st.session_state.edit_mode = edit_mode
                        st.rerun()
                with col_cancel:
                    if st.button(
                        "❌ Cancel", key=f"cancel_{q_idx}", use_container_width=True
                    ):
                        edit_mode = st.session_state.get("edit_mode", {})
                        edit_mode[q_idx] = False
                        st.session_state.edit_mode = edit_mode
                        st.rerun()

            # Answer evaluation section
            ans_key = f"answer_{q_idx}"
            current_answer = st.session_state.get("answers", {}).get(q_idx, "")
            answer = st.text_area(
                "Your answer:",
                value=current_answer,
                key=ans_key,
                height=100,
                help="Type the candidate's answer here to evaluate it",
            )
            if "answers" not in st.session_state:
                st.session_state["answers"] = {}
            st.session_state["answers"][q_idx] = answer

            col_eval, col_score = st.columns([3, 1])
            with col_eval:
                if st.button(
                    "📊 Evaluate Answer",
                    key=f"eval_{q_idx}",
                    use_container_width=True,
                    disabled=not answer.strip(),
                ):
                    evaluate_single_question(
                        question,
                        answer,
                        q_idx,
                    )
            with col_score:
                eval_result = st.session_state.get("evaluations", {}).get(q_idx)
                if eval_result:
                    score = eval_result.get("score", 0)
                    score_color = "🟢" if score >= 4 else "🟡" if score >= 3 else "🔴"
                    st.metric("Score", f"{score_color} {score}/5")

            # Show evaluation details
            eval_result = st.session_state.get("evaluations", {}).get(q_idx)
            if eval_result:
                with st.expander(f"Evaluation details for Q{q_idx}", expanded=False):
                    if eval_result.get("strengths"):
                        st.markdown("**Strengths:**")
                        for s in eval_result["strengths"]:
                            st.markdown(f"- {s}")
                    if eval_result.get("gaps"):
                        st.markdown("**Gaps:**")
                        for g in eval_result["gaps"]:
                            st.markdown(f"- {g}")
                    if eval_result.get("follow_up"):
                        st.markdown("**Follow-up questions:**")
                        for f in eval_result["follow_up"]:
                            st.markdown(f"- {f}")
                    if eval_result.get("feedback"):
                        st.markdown(f"**Feedback:** {eval_result['feedback']}")

            st.markdown("---")

    if "resume_insights" in st.session_state and st.session_state.resume_insights:
        st.markdown("---")
        st.subheader("🔑 Resume Key Insights")
        display_resume_insights(st.session_state.resume_insights)
    if "ats_suggestions" in st.session_state and st.session_state.ats_suggestions:
        st.markdown("---")
        st.subheader("📝 ATS Optimization Suggestions")
        display_ats_suggestions(st.session_state.ats_suggestions)

    # Export options
    st.markdown("---")
    st.subheader("📤 Export Questions")

    col1, col2 = st.columns(2)

    with col1:
        if st.button("📄 Export as CSV", use_container_width=True):
            export_to_csv(filtered_questions)
        if st.button("📋 Export as JSON", use_container_width=True):
            export_to_json(filtered_questions)

    with col2:
        if st.button("📋 Export as PDF", use_container_width=True):
            export_to_pdf(filtered_questions)
        if st.button("📘 Export as Word", use_container_width=True):
            export_to_docx(filtered_questions)
        if st.button("📋 Copy to Clipboard", use_container_width=True):
            export_to_clipboard(filtered_questions)


def export_to_csv(questions: List[Dict[str, Any]]) -> None:
    """Export questions to CSV format"""
    df = pd.DataFrame(questions)
    csv = df.to_csv(index=False)

    st.download_button(
        label="💾 Download CSV",
        data=csv,
        file_name="interview_questions.csv",
        mime="text/csv",
    )


def export_to_clipboard(questions: List[Dict[str, Any]]) -> None:
    """Export questions to clipboard-friendly format"""
    formatted_text = format_questions_for_export(questions)

    st.text_area("Copy the questions below:", formatted_text, height=300)
    st.success("📋 Questions formatted for clipboard!")


def display_resume_insights(insights: Dict[str, Any]) -> None:
    """Display extracted resume insights in a structured format."""
    if not insights:
        st.info("No insights extracted from resume.")
        return
    if "technologies" in insights:
        st.markdown(
            f"**Technologies/Skills:** {', '.join(insights['technologies']) if insights['technologies'] else 'N/A'}"
        )
    if "companies" in insights and insights["companies"]:
        st.markdown("**Companies & Durations:**")
        for c in insights["companies"]:
            st.markdown(f"- {c.get('name', 'Unknown')} ({c.get('duration', 'N/A')})")
    if "total_years_experience" in insights:
        st.markdown(f"**Total Years Experience:** {insights['total_years_experience']}")
    if "education" in insights and insights["education"]:
        st.markdown("**Education:**")
        for edu in insights["education"]:
            st.markdown(
                f"- {edu.get('degree', 'N/A')} at {edu.get('institution', 'N/A')} ({edu.get('year', 'N/A')})"
            )
    if "certifications" in insights and insights["certifications"]:
        st.markdown(f"**Certifications:** {', '.join(insights['certifications'])}")
    if "major_projects" in insights and insights["major_projects"]:
        st.markdown("**Major Projects:**")
        for proj in insights["major_projects"]:
            st.markdown(f"- {proj}")


def display_ats_suggestions(suggestions: List[str]) -> None:
    """Display ATS optimization suggestions."""
    if not suggestions:
        st.info("No ATS suggestions available.")
        return
    st.markdown("**Suggestions to improve ATS compatibility:**")
    for s in suggestions:
        st.markdown(f"- {s}")


def export_to_json(questions: List[Dict[str, Any]]) -> None:
    """Export questions to JSON format"""
    data = json.dumps(questions, indent=2)
    st.download_button(
        label="💾 Download JSON",
        data=data,
        file_name="interview_questions.json",
        mime="application/json",
    )


def export_to_pdf(questions: List[Dict[str, Any]]) -> None:
    """Export questions to PDF format"""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title="Interview Questions",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontSize=18,
        textColor=colors.HexColor("#1f77b4"),
        spaceAfter=20,
    )
    heading_style = ParagraphStyle(
        "CustomHeading",
        parent=styles["Heading2"],
        fontSize=14,
        textColor=colors.HexColor("#1f77b4"),
        spaceBefore=12,
        spaceAfter=6,
    )
    question_style = ParagraphStyle(
        "CustomQuestion",
        parent=styles["Normal"],
        fontSize=11,
        spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "CustomMeta",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#666666"),
        spaceAfter=2,
    )

    story = []
    story.append(Paragraph("Interview Questions", title_style))
    story.append(Spacer(1, 12))

    categories: Dict[str, List[Dict[str, Any]]] = {}
    for q in questions:
        cat = q.get("category", "General")
        if cat not in categories:
            categories[cat] = []
        categories[cat].append(q)

    for cat, cat_questions in categories.items():
        story.append(Paragraph(f"{cat}", heading_style))
        for i, q in enumerate(cat_questions, 1):
            diff = q.get("difficulty", "Unknown")
            q_text = f"<b>{i}.</b> {q.get('question', '')}"
            story.append(Paragraph(q_text, question_style))
            story.append(Paragraph(f"Difficulty: {diff}", meta_style))
            story.append(Spacer(1, 8))

    doc.build(story)
    buffer.seek(0)

    st.download_button(
        label="💾 Download PDF",
        data=buffer,
        file_name="interview_questions.pdf",
        mime="application/pdf",
    )


def export_to_docx(questions: List[Dict[str, Any]]) -> None:
    """Export questions to Word (.docx) format"""
    doc = Document()

    doc.add_heading("Interview Questions", level=1)

    categories: Dict[str, List[Dict[str, Any]]] = {}
    for q in questions:
        cat = q.get("category", "General")
        if cat not in categories:
            categories[cat] = []
        categories[cat].append(q)

    for cat, cat_questions in categories.items():
        doc.add_heading(cat, level=2)
        for i, q in enumerate(cat_questions, 1):
            diff = q.get("difficulty", "Unknown")
            p = doc.add_paragraph()
            p.add_run(f"{i}. ").bold = True
            p.add_run(q.get("question", ""))
            p.add_run(f"\nDifficulty: {diff}").italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    st.download_button(
        label="💾 Download Word",
        data=buffer,
        file_name="interview_questions.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


if __name__ == "__main__":
    main()
